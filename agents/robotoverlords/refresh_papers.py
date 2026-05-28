"""Refresh bundled paper/context. Run manually when sources update.

Usage: uv run python agents/robotoverlords/refresh_papers.py

For PDFs: extracts embedded text via pypdf AND renders each page as a PNG so
the agent can read both text and figures.
"""
import io
import re
import shutil
import urllib.request
from pathlib import Path

import pymupdf
from pptx import Presentation
from pypdf import PdfReader

FIGURE_RE = re.compile(r"\b(Figure|Fig\.|Table|Algorithm)\s+(\d+)", re.IGNORECASE)

REPO_ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = Path(__file__).with_name("papers")
PAGE_RENDER_DPI = 150

PDF_SOURCES = {
    "flow_of_options": "https://arxiv.org/pdf/2502.12929",
}

PPTX_SOURCES = {
    "hackathon_context.txt": REPO_ROOT / "Scientific Discovery Hackathon.pptx",
}


def fetch_pdf_bytes(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "hackathon-agent"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return resp.read()


def _figure_mentions(text: str) -> list[str]:
    seen: dict[str, None] = {}
    for kind, num in FIGURE_RE.findall(text):
        kind = "Fig." if kind.lower() in {"fig.", "figure"} else kind.title()
        seen.setdefault(f"{kind} {num}", None)
    return list(seen)


def extract_pdf_text_and_render(pdf_bytes: bytes, out_dir: Path, dpi: int = PAGE_RENDER_DPI) -> str:
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True)

    reader = PdfReader(io.BytesIO(pdf_bytes))
    page_texts = [page.extract_text() or "" for page in reader.pages]

    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    scale = dpi / 72
    matrix = pymupdf.Matrix(scale, scale)
    n_total = len(doc)
    blocks: list[str] = []
    for i, page in enumerate(doc, 1):
        pix = page.get_pixmap(matrix=matrix)
        (out_dir / f"page_{i:03d}.png").write_bytes(pix.tobytes("png"))

        text = page_texts[i - 1].strip() if i - 1 < len(page_texts) else ""
        mentions = _figure_mentions(text)
        image_count = len(page.get_images(full=False))

        header = f"--- PAGE {i} of {n_total} ---"
        notes: list[str] = []
        if mentions:
            notes.append(f"references: {', '.join(mentions)}")
        if image_count:
            notes.append(f"{image_count} embedded image(s)")
        if notes:
            header += f"  [{'; '.join(notes)} — call read_paper_pages(pages=[{i}]) to view]"
        blocks.append(f"{header}\n{text}")
    doc.close()
    return "\n\n".join(blocks)


def extract_pptx_text(path: Path) -> str:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        parts.append(row_text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Notes] {notes}")
        slides.append("\n".join(parts))
    return "\n\n".join(slides)


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    for base, url in PDF_SOURCES.items():
        print(f"processing {base} from {url}")
        pdf_bytes = fetch_pdf_bytes(url)
        text = extract_pdf_text_and_render(pdf_bytes, OUT_DIR / f"{base}_pages")
        (OUT_DIR / f"{base}.txt").write_text(text)
        n = len(list((OUT_DIR / f"{base}_pages").glob("page_*.png")))
        print(f"  wrote {base}.txt ({len(text)} chars) + {n} page PNGs")
    for filename, path in PPTX_SOURCES.items():
        print(f"processing {filename} from {path.name}")
        text = extract_pptx_text(path)
        (OUT_DIR / filename).write_text(text)
        print(f"  wrote {filename}: {len(text)} chars")


if __name__ == "__main__":
    main()
